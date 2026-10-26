"""Checks on the empty Excel templates of Template/Materials (the format MAGNETISiM imports)."""

import re
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
MATERIALS = ROOT / "Template" / "Materials"
POWER_LOSS = MATERIALS / "PowerLoss" / "100kHz_pv.xlsx"
PERMEABILITY = MATERIALS / "PermeabilityFrequency" / "permeability_vs_frequency.xlsx"

# Measured in the importers of the application (packages/matModeling):
# - corePowerLoss: pd.read_excel(header=0), first sheet, column 0 = B [mT], column 1 = Pv [kW/m3],
#   at least 4 rows, frequency read from the file name with r"(\d+(?:\.\d+)?)\s*khz".
# - corePermeabilityFreq: pd.read_excel(header=1), first sheet, row 1 free, row 2 = headers,
#   columns f [Hz], mu' and mu'', at least 4 rows.
POWER_LOSS_HEADERS = ["BmT", "Pv"]
PERMEABILITY_HEADERS = ["f_Hz", "mu_prime", "mu_doubleprime"]
FILENAME_FREQUENCY = re.compile(r"(\d+(?:\.\d+)?)\s*khz")


def rows(path):
    workbook = openpyxl.load_workbook(path)
    assert len(workbook.sheetnames) == 1
    sheet = workbook.worksheets[0]
    return workbook, [
        [cell.value for cell in row] for row in sheet.iter_rows(max_col=sheet.max_column)
    ]


def only_headers(values):
    return [[v for v in row if v is not None] for row in values if any(v is not None for v in row)]


def test_power_loss_template_has_the_headers_and_no_data():
    workbook, values = rows(POWER_LOSS)
    assert workbook.sheetnames == ["Sheet1"]
    assert only_headers(values) == [POWER_LOSS_HEADERS]
    assert values[0] == POWER_LOSS_HEADERS


def test_power_loss_file_name_gives_the_frequency():
    match = FILENAME_FREQUENCY.search(POWER_LOSS.name.lower())
    assert match and float(match.group(1)) == 100.0


def test_permeability_template_has_a_title_row_the_headers_and_no_data():
    workbook, values = rows(PERMEABILITY)
    assert workbook.sheetnames == ["permeability_fit"]
    non_empty = only_headers(values)
    assert len(non_empty) == 2, "a title row and the header row, nothing else"
    assert len(non_empty[0]) == 1 and isinstance(non_empty[0][0], str)
    assert values[1][:3] == PERMEABILITY_HEADERS
    assert non_empty[1] == PERMEABILITY_HEADERS


def test_no_loose_workbooks_in_template_materials():
    assert list(MATERIALS.glob("*.xlsx")) == []
    found = sorted(path.relative_to(MATERIALS).as_posix() for path in MATERIALS.rglob("*.xlsx"))
    assert found == [
        "PermeabilityFrequency/permeability_vs_frequency.xlsx",
        "PowerLoss/100kHz_pv.xlsx",
    ]


def test_the_guide_links_resolve_to_the_template_folders():
    guide = (ROOT / "docs" / "core-library.md").read_text(encoding="utf-8")
    for folder in ("PowerLoss", "PermeabilityFrequency"):
        assert f"(../Template/Materials/{folder}/)" in guide
        assert (MATERIALS / folder).is_dir()


def test_readme_documents_columns_and_units():
    text = (MATERIALS / "README.md").read_text(encoding="utf-8")
    for needle in (
        "`BmT`",
        "`Pv`",
        "mT",
        "kW/m³",
        "`f_Hz`",
        "`mu_prime`",
        "`mu_doubleprime`",
        "Hz",
    ):
        assert needle in text, needle
    assert "kHz" in text and "Add Permeability(f) Curve" in text and "Power Loss" in text


@pytest.mark.parametrize("path", [POWER_LOSS, PERMEABILITY], ids=lambda p: p.name)
def test_templates_are_binary_in_git(path):
    assert path.read_bytes()[:2] == b"PK"
